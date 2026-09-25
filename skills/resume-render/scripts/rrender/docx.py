"""DOCX writer: the same order and text as the render model.

The name is a bold paragraph at the top of the body, never in a page header.
Section headings use the built-in Heading 1 style and bullets use List Bullet,
so an ATS reads the structure. No tables, text boxes, headers or footers.
Core properties hold only the title and author.
"""
from __future__ import annotations

from pathlib import Path

import docx
from docx.shared import Pt, RGBColor

from .model import DASH, SEP, Document

FONT = "Calibri"
BLACK = RGBColor(0, 0, 0)


def _style(document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    heading = document.styles["Heading 1"]
    heading.font.name = FONT
    heading.font.size = Pt(12)
    heading.font.bold = True
    heading.font.color.rgb = BLACK
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(4)
    bullet = document.styles["List Bullet"]
    bullet.paragraph_format.space_after = Pt(2)


def _properties(document, name: str) -> None:
    props = document.core_properties
    props.title = f"{name} resume"
    props.author = name
    for field in ("subject", "keywords", "comments", "category", "last_modified_by", "content_status"):
        setattr(props, field, "")


def write(doc: Document, path: Path) -> None:
    document = docx.Document()
    _style(document)
    _properties(document, doc.name)

    name = document.add_paragraph().add_run(doc.name)
    name.bold = True
    name.font.size = Pt(20)
    if doc.label:
        document.add_paragraph().add_run(doc.label).font.size = Pt(12)
    if doc.contacts:
        document.add_paragraph(SEP.join(c.text for c in doc.contacts))

    for section in doc.sections:
        document.add_heading(section.heading, level=1)
        if section.paragraph:
            document.add_paragraph(section.paragraph)
        for line in section.lines:
            paragraph = document.add_paragraph()
            if line.label:
                paragraph.add_run(f"{line.label}: ").bold = True
            paragraph.add_run(line.text)
        for i, item in enumerate(section.items):
            if item.title:
                title = document.add_paragraph()
                title.paragraph_format.space_before = Pt(6 if i else 0)
                title.paragraph_format.keep_with_next = True
                title.add_run(item.title).bold = True
            meta = [DASH.join(item.dates)] if item.dates else []
            meta += [d.text for d in item.details]
            if meta:
                document.add_paragraph(SEP.join(meta)).paragraph_format.keep_with_next = bool(item.bullets)
            for bullet in item.bullets:
                document.add_paragraph(bullet, style="List Bullet")
    document.save(str(path))


def text(path) -> str:
    """The DOCX body text, one line per paragraph."""
    return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
