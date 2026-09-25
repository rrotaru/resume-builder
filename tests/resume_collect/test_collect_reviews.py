"""Performance reviews: finding, dating, reading and normalizing the files."""
import importlib.util
import os
from datetime import datetime, timezone

import pytest

from rcollect import reviews
from rcollect.common import Page
from rcore import documents

TEXT = "2024 annual review\n\nJordan led the PAY-42 latency work and\nmentored two engineers.\n"


def write(path, text=TEXT, modified="2025-01-20T12:00:00"):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    moment = datetime.fromisoformat(modified).replace(tzinfo=timezone.utc).timestamp()
    os.utime(path, (moment, moment))
    return path


def test_review_files_recurse_skip_hidden_and_split_unsupported(tmp_path):
    for rel in ("b.txt", "a/c.md", "a/d.PDF", ".hidden.txt", ".git/x.txt", "e.doc", "f.docx"):
        write(tmp_path / rel)
    supported, unsupported = reviews.review_files(tmp_path)
    assert [p.relative_to(tmp_path).as_posix() for p in supported] == ["a/c.md", "a/d.PDF", "b.txt", "f.docx"]
    assert [p.name for p in unsupported] == ["e.doc"]


def test_dates_from_the_name_or_the_modification_time(tmp_path):
    assert reviews.file_date(write(tmp_path / "review 2024-12-15 final.txt")) == ("2024-12-15", "name")
    assert reviews.file_date(write(tmp_path / "2024-02-30 then 2024-03-01.txt")) == ("2024-03-01", "name")
    assert reviews.file_date(write(tmp_path / "2024-H2.txt")) == ("2025-01-20", "modified")
    assert reviews.file_date(write(tmp_path / "12024-12-150.txt")) == ("2025-01-20", "modified")


def test_read_review(tmp_path):
    path = write(tmp_path / "2024" / "annual.md")
    item = reviews.read_review(tmp_path, path)
    assert item["file"] == "2024/annual.md" and item["format"] == "md" and item["path"] == str(path.resolve())
    assert item["text"] == TEXT and item["sha256"].startswith("sha256:")
    assert (item["date"], item["date_from"]) == ("2025-01-20", "modified")


def test_read_review_without_text(tmp_path):
    with pytest.raises(documents.NoText):
        reviews.read_review(tmp_path, write(tmp_path / "short.txt", "Too short."))
    path = tmp_path / "latin1.txt"
    path.write_bytes("Jordan’s review".encode("cp1252") * 10)
    with pytest.raises(documents.ExtractError, match="not UTF-8 text"):
        reviews.read_review(tmp_path, path)


def test_normalize_reviews():
    items = [{"file": "2024/annual.pdf", "text": TEXT, "date": "2025-01-20"},
             {"file": "blank.txt", "text": "\n\n", "date": "2025-01-21"},
             {"file": "", "text": TEXT, "date": "2025-01-20"},
             {"file": "x.txt", "date": "2025-01-20"},
             {"file": "x.txt", "text": TEXT, "date": "Jan 2025"}]
    result = reviews.normalize([Page(1, items)], "01-raw/reviews.jsonl", "01-raw.tmp/reviews.jsonl")
    assert result.bad == ["01-raw.tmp/reviews.jsonl:1#/items/2: no file name",
                          "01-raw.tmp/reviews.jsonl:1#/items/3: no text",
                          "01-raw.tmp/reviews.jsonl:1#/items/4: date 'Jan 2025' is not YYYY-MM-DD"]
    annual, blank = result.drafts
    assert (annual.source, annual.kind, annual.engineer_role, annual.native_key) == (
        "review", "perf_review", "subject", "2024/annual")
    assert annual.title == "2024 annual review"
    assert annual.excerpt == "Jordan led the PAY-42 latency work and mentored two engineers."
    assert annual.created_at == "2025-01-20T00:00:00Z" and annual.url is None
    assert annual.raw_ref == "01-raw/reviews.jsonl:1#/items/0"
    assert ("jira", "PAY-42") in annual.refs
    assert (blank.title, blank.excerpt) == ("", "")


needs_docx = pytest.mark.skipif(
    not os.environ.get("CI") and importlib.util.find_spec("docx") is None,
    reason="needs ingest_reviews.py's dependencies (python-docx)")


@needs_docx
def test_read_docx_review(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("2025 H1 review")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Impact"
    table.cell(0, 1).text = "Jordan shipped the idempotency cache and cut checkout latency."
    path = tmp_path / "2025-07-01 review.docx"
    document.save(path)
    item = reviews.read_review(tmp_path, path)
    assert item["format"] == "docx" and (item["date"], item["date_from"]) == ("2025-07-01", "name")
    assert item["text"].splitlines()[:3] == ["2025 H1 review", "Impact",
                                             "Jordan shipped the idempotency cache and cut checkout latency."]


def test_review_ingest_pins_the_same_versions_as_import():
    """The test command installs extract_text.py's requirements, so ingest_reviews.py must match them."""
    import re

    from conftest import REPO

    def pins(rel):
        text = (REPO / rel).read_text(encoding="utf-8")
        return re.search(r"# dependencies = \[\n(.*?)# \]", text, re.DOTALL).group(1)

    assert pins("skills/resume-collect/scripts/ingest_reviews.py") == pins(
        "skills/resume-import/scripts/extract_text.py")
