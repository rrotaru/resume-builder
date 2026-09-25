"""PDF and DOCX extraction. Needs extract_text.py's dependencies:
uv run --with pytest --with-requirements skills/resume-import/scripts/extract_text.py pytest

They skip when missing, except in CI (CI set), where they fail.
"""
import io
import os

import pytest

if not os.environ.get("CI"):
    for module in ("docx", "pypdf"):
        pytest.importorskip(module, reason="needs extract_text.py's dependencies (see module docstring)")

import docx  # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402
from rimport import extract  # noqa: E402
from samples import encrypt_pdf, make_pdf  # noqa: E402

EXTRA_NS = ('xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
            'xmlns:v="urn:schemas-microsoft-com:vml"')
LINES = ["Jordan Rivera", "Senior Software Engineer, Northwind Payments",
         "Jan 2023 – Present", "Built a Redis-backed idempotency cache (Go)"]


def pdf_text(data, name="resume.pdf"):
    return extract.extract(data, "pdf", name)


def test_pdf_text_pages_and_links():
    data = make_pdf([LINES, ["Page two: Software Engineer, Tailspin Toys"]],
                    links=["https://github.com/jrivera", "mailto:jordan@example.com"])
    result = pdf_text(data)
    assert result.pages == 2
    lines = result.text.splitlines()
    assert lines[:4] == LINES
    assert "Page two: Software Engineer, Tailspin Toys" in lines
    assert lines[-3:] == ["Links in the document:", "https://github.com/jrivera", "mailto:jordan@example.com"]
    assert result.text.endswith("\n")


def test_pdf_with_no_text_is_no_text():
    with pytest.raises(extract.NoText, match=r"no text found in scan.pdf \(1 page\)\. It looks like a scanned PDF"):
        pdf_text(make_pdf([[]]), "scan.pdf")


def test_pdf_with_an_empty_user_password_opens():
    data = encrypt_pdf(make_pdf([LINES]), user_password="")
    assert pdf_text(data).text.splitlines()[:4] == LINES


def test_pdf_with_a_password_is_refused():
    data = encrypt_pdf(make_pdf([LINES]), user_password="secret")
    with pytest.raises(extract.ExtractError, match="resume.pdf is password-protected; save an unprotected copy"):
        pdf_text(data)


def test_damaged_pdf():
    with pytest.raises(extract.ExtractError, match="resume.pdf: cannot be read as a PDF"):
        pdf_text(b"%PDF-1.4\nnot really a pdf")


# DOCX ------------------------------------------------------------------------

def _hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{r_id}"><w:r><w:t>{text}</w:t></w:r></w:hyperlink>')
    paragraph._p.append(link)


def _text_box(paragraph, text):
    """A text box as Word saves it: the shape in mc:Choice, a VML copy in mc:Fallback."""
    run = parse_xml(
        f'<w:r {nsdecls("w")} {EXTRA_NS}><mc:AlternateContent>'
        f'<mc:Choice Requires="wps"><w:drawing><wps:txbx><w:txbxContent>'
        f'<w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent></wps:txbx></w:drawing></mc:Choice>'
        f'<mc:Fallback><w:pict><v:textbox><w:txbxContent>'
        f'<w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent></v:textbox></w:pict></mc:Fallback>'
        f'</mc:AlternateContent></w:r>')
    paragraph._p.append(run)


def make_docx():
    document = docx.Document()
    header = document.sections[0].header
    header.paragraphs[0].text = "Jordan Rivera | jordan@example.com"
    document.sections[0].footer.paragraphs[0].text = "References available on request"
    document.add_paragraph("Senior Software Engineer, Northwind Payments")
    dates = document.add_paragraph("Jan 2023")
    dates.add_run("\t")
    dates.add_run("Present")
    hidden = document.add_paragraph("Visible bullet")
    secret = hidden.add_run(" KEYWORD STUFFING")
    secret.font.hidden = True
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Software Engineer, Tailspin Toys"
    table.cell(0, 1).text = "June 2019 – December 2022"
    links = document.add_paragraph("Profiles: ")
    _hyperlink(links, "https://github.com/jrivera", "GitHub")
    anchor = document.add_paragraph("Projects")
    _text_box(anchor, "ledger-lint, an open-source linter")
    breaks = document.add_paragraph("Line one")
    breaks.add_run().add_break()
    breaks.add_run("Line two")
    field = document.add_paragraph()
    field._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> HYPERLINK '
                              f'"https://jrivera.dev" </w:instrText></w:r>'))
    field.add_run("jrivera.dev")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_docx_text_in_order():
    lines = extract.extract(make_docx(), "docx", "resume.docx").text.splitlines()
    assert lines == [
        "Jordan Rivera | jordan@example.com",  # header first
        "Senior Software Engineer, Northwind Payments",
        "Jan 2023 Present",  # a tab is a space
        "Visible bullet",  # hidden text skipped
        "Software Engineer, Tailspin Toys",  # table cells in order
        "June 2019 – December 2022",
        "Profiles: GitHub",
        "Projects",
        "ledger-lint, an open-source linter",  # the text box once, not its fallback copy
        "Line one",
        "Line two",
        "jrivera.dev",
        "References available on request",  # footer last
        "",
        "Links in the document:",
        "https://github.com/jrivera",
        "https://jrivera.dev",
    ]


def test_docx_hidden_text_can_be_switched_off():
    document = docx.Document(io.BytesIO(make_docx()))
    for vanish in document.element.body.iter(qn("w:vanish")):
        vanish.set(qn("w:val"), "false")
    out = io.BytesIO()
    document.save(out)
    assert "Visible bullet KEYWORD STUFFING" in extract.extract(out.getvalue(), "docx", "r.docx").text


def test_empty_docx_is_no_text():
    out = io.BytesIO()
    docx.Document().save(out)
    with pytest.raises(extract.NoText, match=r"^no text found in r.docx\. Provide another file"):
        extract.extract(out.getvalue(), "docx", "r.docx")


def test_damaged_docx():
    with pytest.raises(extract.ExtractError, match="r.docx: cannot be read as a DOCX"):
        extract.extract(b"PK\x03\x04 not a zip", "docx", "r.docx")


# extract_text.py with a PDF ------------------------------------------------------

def test_extract_text_cli_pdf(workspace, tmp_path, capsys):
    import extract_text
    from rcore import wsio

    resume = tmp_path / "resume.pdf"
    resume.write_bytes(make_pdf([LINES], links=["https://github.com/jrivera"]))
    assert extract_text.main(["--workspace", str(workspace), "--resume", str(resume)]) == 0
    out = capsys.readouterr().out
    assert f"extracted {resume} (pdf, 1 page): 7 lines in 03-profile.tmp/resume.txt" in out
    assert "   3  Jan 2023 – Present\n" in out
    assert "   7  https://github.com/jrivera\n" in out
    assert wsio.read_json(workspace / "03-profile.tmp" / "source.json")["format"] == "pdf"


def test_extract_text_cli_scanned_pdf(workspace, tmp_path, capsys):
    import extract_text
    from rcore import wsio

    before = wsio.read_json(workspace / "config.json")
    scan = tmp_path / "scan.pdf"
    scan.write_bytes(make_pdf([[], []]))
    assert extract_text.main(["--workspace", str(workspace), "--resume", str(scan)]) == 3
    assert capsys.readouterr().err == (
        f"error: no text found in {scan} (2 pages). It looks like a scanned PDF, and OCR is not supported. "
        "Provide a DOCX or TXT version, or paste the text.\n")
    assert wsio.read_json(workspace / "config.json") == before
    assert not (workspace / "03-profile.tmp").exists()
