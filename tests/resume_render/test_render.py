"""render.py end to end. Needs render.py's dependencies:
uv run --with pytest --with-requirements skills/resume-render/scripts/render.py pytest

Tests marked chromium also need Playwright's Chromium. They skip when it is
missing, except in CI (CI set), where they fail.
"""
import json
import os
import subprocess
import sys

import pytest
from conftest import REPO

if not os.environ.get("CI"):
    for module in ("docx", "jinja2", "playwright", "pypdf"):
        pytest.importorskip(module, reason="needs render.py's dependencies (see module docstring)")

import docx as pydocx  # noqa: E402
import render  # noqa: E402
from pypdf import PdfReader  # noqa: E402
from rcore import stages, wsio  # noqa: E402
from rrender import model, outcheck, pdf  # noqa: E402

SCRIPT = REPO / "skills" / "resume-render" / "scripts" / "render.py"
SAVED = REPO / "tests" / "fixtures" / "render"
GENERAL = "08-ats/general/resume.json"
JOB = "08-ats/jobs/fintech-sre/resume.json"


def run(workspace, *args):
    return render.main(["--workspace", str(workspace), *args])


def needs_chromium(code):
    if code == render.NO_CHROMIUM:
        if os.environ.get("CI"):
            pytest.fail("Chromium is not installed in CI")
        pytest.skip("Chromium is not installed; run render.py --install-browser")
    return code


def edit(workspace, rel, change):
    data = wsio.read_json(workspace / rel)
    change(data)
    wsio.write_json(workspace / rel, data)


def set_bullet(resume, work, index, text):
    resume["work"][work]["x-highlights"][index]["text"] = text
    resume["work"][work]["highlights"][index] = text


def document(workspace, rel):
    return model.build(wsio.read_json(workspace / rel))


def test_check_writes_nothing(workspace, capsys):
    assert run(workspace, "--check") == 0
    assert capsys.readouterr().out == "render check passed\n"
    assert not (workspace / "out").exists()


def test_no_pdf_writes_docx_and_txt(workspace):
    assert run(workspace, "--no-pdf") == 0
    out = workspace / "out"
    assert sorted(p.relative_to(out).as_posix() for p in out.rglob("*") if p.is_file()) == [
        "_stage.json", "general/resume.docx", "general/resume.txt",
        "jobs/fintech-sre/resume.docx", "jobs/fintech-sre/resume.txt", "stories.md"]
    assert (out / "general" / "resume.txt").read_bytes() == (SAVED / "general.txt").read_bytes()
    assert (out / "jobs" / "fintech-sre" / "resume.txt").read_bytes() == (SAVED / "job-fintech-sre.txt").read_bytes()
    assert (out / "stories.md").read_bytes() == (workspace / "07-sanitized" / "stories.md").read_bytes()
    meta = wsio.read_json(out / "_stage.json")
    assert meta["extra"] == {"targets": ["general", "fintech-sre"], "pdf": False, "pages": {}, "stories": True}
    assert set(meta["inputs"]) >= {GENERAL, JOB, "decisions/terms.json", "07-sanitized/stories.md"}
    assert stages.status(workspace)["out"] == "fresh"
    assert not (workspace / "out.tmp").exists()


def test_docx_structure(workspace):
    assert run(workspace, "--no-pdf") == 0
    word = pydocx.Document(str(workspace / "out" / "general" / "resume.docx"))
    headings = [p.text for p in word.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["Experience", "Projects", "Skills", "Education", "Certifications"]
    bullets = [p.text for p in word.paragraphs if p.style.name == "List Bullet"]
    doc = document(workspace, GENERAL)
    assert bullets == [b for s in doc.sections for item in s.items for b in item.bullets]
    assert word.paragraphs[0].text == "Jordan Rivera" and not word.tables
    assert all(not p.text for s in word.sections for p in [*s.header.paragraphs, *s.footer.paragraphs])
    props = word.core_properties
    assert (props.title, props.author, props.comments, props.keywords) == (
        "Jordan Rivera resume", "Jordan Rivera", "", "")


def test_failed_gate_writes_nothing_and_keeps_the_previous_output(workspace, capsys):
    assert run(workspace, "--no-pdf") == 0
    before = stages.hash_path(workspace / "out")
    edit(workspace, GENERAL, lambda r: r.update(awards=[{"title": "Engineer of the year"}]))
    capsys.readouterr()
    assert run(workspace, "--no-pdf") == 1
    assert capsys.readouterr().out == (
        "[general]\n08-ats/general/resume.json: $.awards: unexpected property\n  fix: /resume-builder:ats\n")
    assert stages.hash_path(workspace / "out") == before
    assert not (workspace / "out.tmp").exists()


def test_target_narrows_the_output(workspace):
    assert run(workspace, "--no-pdf", "--target", "general") == 0
    assert sorted(p.name for p in (workspace / "out").iterdir()) == ["_stage.json", "general", "stories.md"]
    assert wsio.read_json(workspace / "out" / "_stage.json")["extra"]["targets"] == ["general"]


def test_stories_are_never_copied_from_06_bullets(workspace, capsys):
    (workspace / "07-sanitized" / "stories.md").unlink()
    assert run(workspace, "--no-pdf") == 0
    assert not (workspace / "out" / "stories.md").exists()
    assert "06-bullets/stories.md is never copied" in capsys.readouterr().out
    assert wsio.read_json(workspace / "out" / "_stage.json")["extra"]["stories"] is False


def test_nothing_to_render(tmp_path, capsys):
    assert run(tmp_path, "--no-pdf") == 1
    assert capsys.readouterr().out == "nothing to render; run /resume-builder:ats\n"


def test_inputs_changed_during_render_writes_nothing(workspace, monkeypatch, capsys):
    real = render.txt.render

    def render_and_edit(doc):
        edit(workspace, "decisions/profile.json", lambda p: p["basics"].update(phone="555-0100"))
        return real(doc)

    monkeypatch.setattr(render.txt, "render", render_and_edit)
    assert run(workspace, "--no-pdf") == 1
    assert "inputs changed during render" in capsys.readouterr().err
    assert not (workspace / "out").exists() and not (workspace / "out.tmp").exists()


def test_stories_that_appear_during_render_are_not_copied(workspace, monkeypatch, capsys):
    stories = workspace / "07-sanitized" / "stories.md"
    stories.unlink()
    real = render.txt.render

    def render_and_add_stories(doc):
        stories.write_text("Led Project Falcon.\n", encoding="utf-8")  # never checked by the gate
        return real(doc)

    monkeypatch.setattr(render.txt, "render", render_and_add_stories)
    assert run(workspace, "--no-pdf") == 1
    assert "inputs changed during render" in capsys.readouterr().err
    assert not (workspace / "out").exists() and not (workspace / "out.tmp").exists()


def test_stories_that_vanish_during_render_abort_it(workspace, monkeypatch, capsys):
    real = render.txt.render

    def render_and_remove_stories(doc):
        (workspace / "07-sanitized" / "stories.md").unlink(missing_ok=True)
        return real(doc)

    monkeypatch.setattr(render.txt, "render", render_and_remove_stories)
    assert run(workspace, "--no-pdf") == 1
    assert "inputs changed during render" in capsys.readouterr().err
    assert not (workspace / "out").exists() and not (workspace / "out.tmp").exists()


def test_output_check_catches_a_term_joined_by_a_control_character(workspace, capsys):
    # The terms check does not treat U+0001 as a separator, but the render model
    # turns it into a space, so the denied term appears in the output.
    edit(workspace, GENERAL, lambda r: set_bullet(r, 1, 0, "Built Project\u0001Falcon in Go"))
    assert run(workspace, "--check") == 0
    assert run(workspace, "--no-pdf") == 1
    out = capsys.readouterr().out
    assert "out/general/resume.txt:14: contains denylisted term 'Project Falcon'" in out
    assert "out/general/resume.docx:" in out
    assert not (workspace / "out").exists()


def test_output_check_reports_missing_text(workspace):
    doc = document(workspace, GENERAL)
    text = (SAVED / "general.txt").read_text(encoding="utf-8").replace("Mentored", "Coached")
    patterns = outcheck.terms.compile_terms(["Project Falcon"])
    assert outcheck.check(doc, {"x.txt": text}, patterns) == [
        "x.txt: missing 'Mentored two new engineers through on-call onboarding'"]
    assert outcheck.missing(doc, "\n".join(reversed(doc.checked_texts()))) != []  # order matters


def test_install_browser_uses_this_environment(monkeypatch):
    calls = []
    monkeypatch.setattr(render.subprocess, "call", lambda command: calls.append(command) or 0)
    assert render.main(["--install-browser", "--with-deps"]) == 0
    assert calls == [[sys.executable, "-m", "playwright", "install", "--with-deps", "chromium"]]


def test_missing_chromium_exits_3_and_writes_nothing(workspace, tmp_path):
    env = dict(os.environ, PLAYWRIGHT_BROWSERS_PATH=str(tmp_path / "no-browsers"))
    result = subprocess.run([sys.executable, str(SCRIPT), "--workspace", str(workspace)],
                            capture_output=True, text=True, env=env)
    assert result.returncode == 3, result.stderr
    assert "render.py --install-browser" in result.stderr and "--no-pdf" in result.stderr
    assert not (workspace / "out").exists() and not (workspace / "out.tmp").exists()


@pytest.mark.chromium
def test_pdf_holds_every_heading_and_bullet(workspace, capsys):
    assert needs_chromium(run(workspace)) == 0
    assert "rendered general: out/general/ resume.pdf (1 page), resume.docx, resume.txt" in capsys.readouterr().out
    for rel, folder in [(GENERAL, "general"), (JOB, "jobs/fintech-sre")]:
        path = workspace / "out" / folder / "resume.pdf"
        reader = PdfReader(str(path))
        assert outcheck.missing(document(workspace, rel), pdf.text(path)) == []
        assert dict(reader.metadata) == {"/Title": "Jordan Rivera resume", "/Author": "Jordan Rivera"}
        fonts = {str(f.get_object()["/BaseFont"]) for page in reader.pages
                 for f in page["/Resources"]["/Font"].get_object().values()}
        assert fonts and all("Carlito" in f for f in fonts)
        assert round(float(reader.pages[0].mediabox.width)) == 612  # US Letter
    meta = wsio.read_json(workspace / "out" / "_stage.json")
    assert meta["extra"]["pdf"] is True and meta["extra"]["pages"] == {"general": 1, "fintech-sre": 1}


@pytest.mark.chromium
def test_a4_paper(workspace):
    assert needs_chromium(run(workspace, "--paper", "a4", "--target", "general")) == 0
    reader = PdfReader(str(workspace / "out" / "general" / "resume.pdf"))
    assert abs(float(reader.pages[0].mediabox.width) - 595.3) < 1  # A4 is 595.3 pt; Chromium uses 8.27 in


@pytest.mark.chromium
def test_markup_in_a_bullet_renders_as_text(workspace):
    text = "Shipped <script>alert(1)</script> & <b>bold</b> fixes in Go"
    edit(workspace, GENERAL, lambda r: set_bullet(r, 1, 0, text))
    assert needs_chromium(run(workspace, "--target", "general")) == 0
    extracted = pdf.text(workspace / "out" / "general" / "resume.pdf")
    assert text in " ".join(extracted.split())


@pytest.mark.chromium
def test_the_page_cannot_fetch_anything():
    try:
        with pdf.Chromium() as chromium:
            html = ('<html><head><link rel="stylesheet" href="https://example.com/x.css"></head>'
                    '<body><img src="http://example.com/x.png">Hello</body></html>')
            data, blocked = chromium.print_pdf(html, "letter")
    except pdf.ChromiumMissing:
        needs_chromium(render.NO_CHROMIUM)
    assert data.startswith(b"%PDF")
    assert sorted(blocked) == ["http://example.com/x.png", "https://example.com/x.css"]
